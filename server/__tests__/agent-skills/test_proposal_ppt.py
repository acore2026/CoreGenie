"""Real source-to-PPT tests, plus isolated rendering and review failure tests."""
import copy
import hashlib
import io
import json
from pathlib import Path
import zipfile
from argparse import Namespace

from PIL import Image
from pptx import Presentation
import pytest

from test_skill_regressions import REVIEW, load, make_docx, paragraph

PPT = load("3gpp-proposal-ppt", "proposal_ppt.py")
QA = load("3gpp-proposal-ppt", "render_qa.py")


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    return {"path": str(path), "sha256": PPT.digest(path)}


@pytest.fixture
def cards(tmp_path):
    source = tmp_path / "S2-2600001.docx"
    drawing = '<w:drawing xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:blip r:embed="rId1"/></w:drawing>'
    fallback = '<w:pict xmlns:v="urn:schemas-microsoft-com:vml"><v:imagedata r:id="rId1"/></w:pict>'
    body = paragraph("Example proposes a 20 ms delay with 10 samples.")
    body += '<w:p><w:r>' + drawing + fallback + '</w:r></w:p>'
    body += '<w:tbl><w:tr><w:tc>' + paragraph("Delay 20 ms") + '<w:p><w:r>' + drawing + '</w:r></w:p></w:tc></w:tr></w:tbl>'
    make_docx(source, body)
    with zipfile.ZipFile(source) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    image = io.BytesIO()
    Image.new("RGB", (400, 100), "blue").save(image, format="PNG")
    parts["word/media/image.png"] = image.getvalue()
    with zipfile.ZipFile(source, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    texts = tmp_path / "texts"
    texts.mkdir()
    code, result = REVIEW.extract_docx(source, texts, tmp_path / "figures")
    assert not result.startswith("ERROR")
    source_index = texts / f"{code}.source.json"
    evidence = [{"code": code, "locator": "BLOCK: 1 p", "quote": "20 ms delay with 10 samples"}]
    claim = {"text": "该提案建议使用 20 ms 延迟，样本数为 10。", "evidence": evidence}
    return {
        "schema": PPT.SCHEMA, "title": "提案技术讨论", "meeting": "合成资料，不代表真实会议结果",
        "sources": [{"path": str(source_index), "sha256": PPT.digest(source_index)}],
        "proposals": [{"code": code, "company": "Example", "title": "Delay proposal", "kind": "贡献稿", "spec": "原文未注明", "clauses": [],
                       "summary": copy.deepcopy(claim), "points": [{"title": "延迟与样本", "claims": [copy.deepcopy(claim)],
                       "image": {"id": "word/media/image.png", "locator": "BLOCK: 2 p", "caption": "合成原图", "readability": "readable"},
                       "table": {"headers": ["指标", "值"], "rows": [["延迟", "20 ms"]], "evidence": [{"code": code, "locator": "BLOCK: 3 tbl", "quote": "Delay 20 ms"}]}}],
                       "watch": [], "dataChecks": [{"value": "20", "unit": "ms", "condition": "10 samples", "evidence": evidence, "status": "verified"}]}],
        "synthesis": [], "failures": [],
    }


def test_source_index_maps_table_images_and_deduplicates_fallback(cards):
    source = json.loads(Path(cards["sources"][0]["path"]).read_text())
    image = source["images"][0]
    assert image["locators"] == ["BLOCK: 2 p", "BLOCK: 3 tbl"]
    assert PPT.digest(image["path"]) == image["sha256"]
    assert Path(source["text"]["path"]).read_text().count("[FIGURE REF:") == 2


def test_extraction_accepts_single_file_and_checkpoints_failed_batch(cards, tmp_path):
    source = json.loads(Path(cards["sources"][0]["path"]).read_text())
    args = Namespace(input=source["source"]["path"], texts=str(tmp_path / "single-texts"), figures=str(tmp_path / "single-figures"))
    REVIEW.cmd_extract(args)
    assert (Path(args.texts) / "S2-2600001.source.json").is_file()
    (tmp_path / "S2-2600002.DOCX").write_bytes(b"corrupt docx")
    (tmp_path / "~$locked.docx").write_bytes(b"lock file")
    args.input = str(tmp_path)
    args.texts, args.figures = str(tmp_path / "batch-texts"), str(tmp_path / "batch-figures")
    with pytest.raises(SystemExit, match="1 extraction"):
        REVIEW.cmd_extract(args)
    summary = json.loads((Path(args.texts) / "extraction-summary.json").read_text())
    assert summary["status"] == "partial"
    assert [r["status"] for r in summary["results"]] == ["extracted", "failed"]
    assert (Path(args.texts) / "S2-2600001.source.json").is_file()


def test_real_ppt_preserves_editable_tables_image_ratio_notes_and_original(cards, tmp_path):
    source = json.loads(Path(cards["sources"][0]["path"]).read_text())
    original = Path(source["source"]["path"]).read_bytes()
    result = PPT.build(cards, tmp_path / "deck")
    deck = Presentation(result["deck"]["path"])
    assert result["status"] == "draft"
    assert result["visualReview"] == "not_reviewed"
    tables = [shape.table for s in deck.slides for shape in s.shapes if shape.has_table]
    assert tables[0].cell(1, 1).text == "20 ms"
    pics = [shape for s in deck.slides for shape in s.shapes if shape.shape_type == 13]
    assert len(pics) == 1 and abs(pics[0].width / pics[0].height - 4) < .001
    assert pics[0].image.blob == Path(source["images"][0]["path"]).read_bytes()
    assert any("20 ms delay with 10 samples" in s.notes_slide.notes_text_frame.text for s in deck.slides)
    assert all("/workspace/" not in s.notes_slide.notes_text_frame.text for s in deck.slides)
    assert Path(source["source"]["path"]).read_bytes() == original
    assert json.loads((tmp_path / "deck/cards.json").read_text()) == cards
    with pytest.raises(ValueError, match="new output"):
        PPT.build(cards, tmp_path / "deck")


def test_synthesis_table_is_editable_and_preserves_all_rows(cards, tmp_path):
    claim = copy.deepcopy(cards["proposals"][0]["summary"])
    claim["table"] = {"headers": ["提案", "值"], "rows": [[f"比较项 {n}", str(n)] for n in range(14)], "evidence": copy.deepcopy(claim["evidence"])}
    cards["synthesis"] = [claim]
    result = PPT.build(cards, tmp_path / "deck")
    deck = Presentation(result["deck"]["path"])
    rows = [row.cells[0].text for slide in deck.slides for shape in slide.shapes if shape.has_table for row in list(shape.table.rows)[1:]]
    assert all(rows.count(f"比较项 {n}") == 1 for n in range(14))
    # Only structure and references are checked here; these synthetic table
    # numbers are deliberately not a claim of semantic verification.


def test_omitted_source_is_rejected(cards, tmp_path):
    source = json.loads(Path(cards["sources"][0]["path"]).read_text())
    source["document"] = "S2-2600002"
    cards["sources"].append(save_json(tmp_path / "second.source.json", source))
    with pytest.raises(ValueError, match="Source index omitted"):
        PPT.validate_cards(cards)
    cards["failures"].append({"code": "S2-2600002", "reason": "尚未分析"})
    assert len(PPT.validate_cards(cards)) == 2


@pytest.mark.parametrize("mutation, message", [
    (lambda c: c["proposals"][0]["summary"].update(evidence=[]), "Evidence"),
    (lambda c: c["proposals"][0]["summary"]["evidence"][0].update(quote="invented claim"), "Quote not found"),
    (lambda c: c["proposals"][0]["summary"]["evidence"][0].update(locator="BLOCK: 99 p"), "Quote not found"),
    (lambda c: c["proposals"][0]["points"][0]["image"].update(locator="BLOCK: 1 p"), "Image locator"),
    (lambda c: c["proposals"][0]["points"][0]["table"]["rows"].append(["bad"]), "columns"),
    (lambda c: c["proposals"][0]["dataChecks"][0].update(value="999"), "Verified value"),
    (lambda c: c["proposals"].append(copy.deepcopy(c["proposals"][0])), "duplicate"),
])
def test_invalid_cards_fail_before_creating_output(cards, tmp_path, mutation, message):
    mutation(cards)
    with pytest.raises(ValueError, match=message):
        PPT.build(cards, tmp_path / "rejected")
    assert not (tmp_path / "rejected").exists()


def test_changed_source_rejected(cards, tmp_path):
    source = json.loads(Path(cards["sources"][0]["path"]).read_text())
    Path(source["text"]["path"]).write_text("different revision")
    with pytest.raises(ValueError, match="File changed"):
        PPT.build(cards, tmp_path / "rejected")


def test_unreadable_figure_and_failed_input_are_visible(cards, tmp_path):
    cards["proposals"][0]["points"][0]["image"]["readability"] = "unreadable"
    cards["failures"] = [{"code": "S2-2600002", "reason": "用户给出的文档损坏，无法提取。"}]
    result = PPT.build(cards, tmp_path / "deck")
    contents = "\n".join(shape.text for s in Presentation(result["deck"]["path"]).slides for shape in s.shapes if shape.has_text_frame)
    assert "原图无法辨认" in contents
    assert "未完成资料" in contents and "文档损坏" in contents


def test_layout_rejects_nonfinite_overlap_and_table_in_notes_area():
    box = dict(slide=0, x=.5, y=2, w=2, h=1, zone="body", label="table")
    assert PPT.audit_boxes([dict(box, x=float("nan"))])
    assert PPT.audit_boxes([dict(box, y=6.2)])
    assert PPT.audit_boxes([box, dict(box, label="text")])
    assert not PPT.audit_boxes([box, dict(box, x=4)])


def test_long_text_fails_without_silently_trimming(cards, tmp_path):
    cards["proposals"][0]["summary"]["text"] = "一" * 500
    with pytest.raises(ValueError, match="more room"):
        PPT.build(cards, tmp_path / "rejected")
    assert not (tmp_path / "rejected").exists()


def test_missing_render_dependencies_leave_explicit_draft_result(cards, tmp_path, monkeypatch):
    deck = PPT.build(cards, tmp_path / "deck")
    monkeypatch.setattr(QA.shutil, "which", lambda name: None)
    result = QA.render(deck["deck"]["path"], tmp_path / "render")
    assert result["status"] == "unavailable" and result["visualReview"] == "not_reviewed"
    assert (tmp_path / "render/render-summary.json").is_file()
    assert Path(deck["deck"]["path"]).is_file()


def test_renderer_checks_output_instead_of_trusting_zero_exit(cards, tmp_path, monkeypatch):
    deck = PPT.build(cards, tmp_path / "deck")
    monkeypatch.setattr(QA.shutil, "which", lambda name: f"/fake/{name}")
    monkeypatch.setattr(QA, "run", lambda command: None)
    result = QA.render(deck["deck"]["path"], tmp_path / "render")
    assert result["status"] == "failed" and "no PDF" in result["error"]
    assert result["visualReview"] == "not_reviewed"


def test_renderer_does_not_accept_missing_pages(cards, tmp_path, monkeypatch):
    deck = PPT.build(cards, tmp_path / "deck")
    output = tmp_path / "render"
    monkeypatch.setattr(QA.shutil, "which", lambda name: f"/fake/{name}")
    def fake_run(command):
        if Path(command[0]).name == "libreoffice":
            (output / "proposals.pdf").write_bytes(b"synthetic pdf")
        else:
            Image.new("RGB", (160, 90), "white").save(output / "slide-1.png")
    monkeypatch.setattr(QA, "run", fake_run)
    result = QA.render(deck["deck"]["path"], output)
    assert result["status"] == "failed" and "rendered pages" in result["error"]
    assert result["visualReview"] == "not_reviewed"


def review_fixture(cards, tmp_path):
    deck = PPT.build(cards, tmp_path / "deck")
    images = []
    for number in range(1, deck["slides"] + 1):
        path = tmp_path / f"slide-{number}.png"
        Image.new("RGB", (160, 90), "white").save(path)
        images.append({"number": number, "path": str(path), "sha256": PPT.digest(path)})
    render_path, review_path = tmp_path / "render.json", tmp_path / "review.json"
    save_json(render_path, {"schema": "3gpp-proposal-render/v1", "kind": "deck", "status": "rendered", "source": deck["deck"], "images": images})
    review = {"deckSha256": deck["deck"]["sha256"], "reviewer": "synthetic test reviewer", "slides": [
        {"number": image["number"], "imageSha256": image["sha256"], "checks": {"layout": "passed", "figures": "not_applicable", "numbers": "not_applicable", "wording": "passed"}, "findings": []} for image in images]}
    save_json(review_path, review)
    return deck, render_path, review_path, review, images


def test_review_binds_all_pages_and_changed_deck_invalidates_it(cards, tmp_path):
    deck, render_path, review_path, review, images = review_fixture(cards, tmp_path)
    assert QA.verify_review(render_path, review_path, tmp_path / "receipt.json")["status"] == "reviewed"
    # A synthetically supplied review exercises the receipt binding, not visual judgment.
    with Path(deck["deck"]["path"]).open("ab") as handle:
        handle.write(b"changed")
    with pytest.raises(ValueError, match="Deck changed"):
        QA.verify_review(render_path, review_path, tmp_path / "stale.json")
    assert not (tmp_path / "stale.json").exists()


@pytest.mark.parametrize("change", ["missing_page", "changed_image", "unresolved", "missing_check"])
def test_incomplete_visual_review_cannot_pass(cards, tmp_path, change):
    _, render_path, review_path, review, images = review_fixture(cards, tmp_path)
    if change == "missing_page":
        review["slides"].pop()
    elif change == "changed_image":
        Path(images[0]["path"]).write_bytes(b"changed")
    elif change == "unresolved":
        review["slides"][0]["findings"] = ["文字被遮挡"]
    else:
        del review["slides"][0]["checks"]["wording"]
    save_json(review_path, review)
    with pytest.raises(ValueError):
        QA.verify_review(render_path, review_path, tmp_path / "rejected.json")
    assert not (tmp_path / "rejected.json").exists()
