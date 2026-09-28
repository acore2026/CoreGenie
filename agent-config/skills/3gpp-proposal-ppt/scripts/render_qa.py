#!/usr/bin/env python3
"""Render local decks/previews and bind a separate visual review to exact files."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    with Path(path).open("x", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def run(command):
    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    if result.returncode:
        raise ValueError(f"{Path(command[0]).name}: {result.stderr or result.stdout}")
    return result


def render(source, output, preview=False):
    source, output = Path(source).resolve(), Path(output).resolve()
    if not source.is_file():
        raise ValueError(f"Input file missing: {source}")
    if source.suffix.lower() not in ({".wmf", ".emf"} if preview else {".pptx"}):
        raise ValueError("Expected WMF/EMF preview input or PPTX render input")
    output.mkdir(parents=True, exist_ok=False)
    report = {"schema": "3gpp-proposal-render/v1", "kind": "preview" if preview else "deck",
              "source": {"path": str(source), "sha256": digest(source)},
              "status": "failed", "visualReview": "not_reviewed", "images": []}
    try:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        rasterizer = shutil.which("pdftoppm")
        if not office or (not preview and not rasterizer):
            report["status"] = "unavailable"
            raise ValueError("LibreOffice and, for PPTX, pdftoppm are required; rebuild the configured sandbox image")
        with tempfile.TemporaryDirectory(prefix="proposal-office-") as profile:
            run([office, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to",
                 "png" if preview else "pdf", "--outdir", str(output), str(source)])
        if preview:
            images = [output / f"{source.stem}.png"]
            expected = 1
        else:
            pdf = output / f"{source.stem}.pdf"
            if not pdf.is_file() or not pdf.stat().st_size:
                raise ValueError("LibreOffice produced no PDF; check that Impress is installed")
            report["pdf"] = {"path": str(pdf), "sha256": digest(pdf)}
            run([rasterizer, "-png", "-scale-to", "1600", str(pdf), str(output / "slide")])
            images = sorted(output.glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
            expected = len(Presentation(source).slides)
        if len(images) != expected or expected == 0:
            raise ValueError(f"Expected {expected} rendered pages, got {len(images)}")
        for n, path in enumerate(images, 1):
            with Image.open(path) as image:
                size = image.size
                image.verify()
            report["images"].append({"number": n, "path": str(path), "sha256": digest(path), "width": size[0], "height": size[1]})
        if digest(source) != report["source"]["sha256"]:
            raise ValueError("Source changed during rendering")
        report["status"] = "rendered"
    except (OSError, ValueError, subprocess.SubprocessError) as exc:
        report["error"] = str(exc)
    write_json(output / "render-summary.json", report)
    return report


def verify_review(render_path, review_path, output):
    render_data = json.loads(Path(render_path).read_text(encoding="utf-8"))
    review = json.loads(Path(review_path).read_text(encoding="utf-8"))
    if render_data.get("schema") != "3gpp-proposal-render/v1" or render_data.get("kind") != "deck" or render_data.get("status") != "rendered":
        raise ValueError("A successful deck render is required")
    source = render_data["source"]
    if digest(source["path"]) != source["sha256"] or review.get("deckSha256") != source["sha256"]:
        raise ValueError("Deck changed; render and review it again")
    expected = len(Presentation(source["path"]).slides)
    images, slides = render_data["images"], review.get("slides", [])
    if len(images) != expected or [i["number"] for i in images] != list(range(1, expected + 1)):
        raise ValueError("Rendered page coverage is incomplete")
    if not isinstance(slides, list) or sorted(s.get("number", -1) for s in slides) != list(range(1, expected + 1)):
        raise ValueError("Review every rendered page exactly once")
    if not isinstance(review.get("reviewer"), str) or not review["reviewer"].strip():
        raise ValueError("Name the reviewer or reviewing Agent")
    for image in images:
        entry = next(s for s in slides if s["number"] == image["number"])
        if digest(image["path"]) != image["sha256"] or entry.get("imageSha256") != image["sha256"]:
            raise ValueError("Rendered image changed; repeat visual review")
        checks = entry.get("checks", {})
        if set(checks) != {"layout", "figures", "numbers", "wording"} or any(value not in {"passed", "not_applicable"} for value in checks.values()) or checks["layout"] != "passed" or checks["wording"] != "passed":
            raise ValueError("Visual/content review has missing checks or unresolved findings")
        if entry.get("findings") != []:
            raise ValueError("Resolve findings before marking the review passed")
    receipt = {"schema": "3gpp-proposal-review/v1", "status": "reviewed",
               "deck": source, "renderSummarySha256": digest(render_path), "reviewSha256": digest(review_path),
               "reviewer": review["reviewer"], "slides": expected,
               "note": "此记录绑定复检者的判断和文件版本，不独立证明内容正确。"}
    write_json(output, receipt)
    return receipt


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    for name in ("render", "preview"):
        command = sub.add_parser(name)
        command.add_argument("--input", required=True)
        command.add_argument("--output", required=True, help="New output directory")
    command = sub.add_parser("verify-review")
    command.add_argument("--render-summary", required=True)
    command.add_argument("--review", required=True)
    command.add_argument("--output", required=True, help="New receipt file")
    args = parser.parse_args()
    if args.command == "verify-review":
        result = verify_review(args.render_summary, args.review, args.output)
    else:
        result = render(args.input, args.output, preview=args.command == "preview")
    print(json.dumps(result, ensure_ascii=False))
    if result["status"] in {"failed", "unavailable"}:
        raise SystemExit(2)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, KeyError, TypeError, OSError) as exc:
        raise SystemExit(f"PPT check failed: {exc}") from exc
