#!/usr/bin/env python3
"""Validate source-linked proposal cards and build an editable, local PPTX."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SCHEMA = "3gpp-proposal-cards/v1"
W, H = 13.333, 7.5


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def text(value, label, limit=2000):
    require(isinstance(value, str) and bool(value.strip()), f"{label}: nonempty text required")
    require(len(value) <= limit, f"{label}: too long; split into separate points/pages")
    require(not any(ord(c) < 32 and c not in "\n\t" for c in value), f"{label}: invalid control character")
    return value


def verified_file(record):
    require(isinstance(record, dict), "File record required")
    path = Path(text(record.get("path"), "file path"))
    require(path.is_absolute() and path.is_file(), f"File not found or not absolute: {path}")
    require(path.stat().st_size <= 250 * 1024 * 1024, f"File exceeds size limit: {path}")
    require(record.get("sha256") == digest(path), f"File changed: {path}")
    return path


def validate_cards(data):
    require(isinstance(data, dict) and data.get("schema") == SCHEMA, "Unsupported cards schema")
    text(data.get("title"), "deck title", 100)
    text(data.get("meeting"), "meeting", 100)
    require(isinstance(data.get("sources"), list) and data["sources"], "Source indexes required")
    sources = {}
    for record in data["sources"]:
        source = json.loads(verified_file(record).read_text(encoding="utf-8"))
        require(source.get("schema") == "3gpp-proposal-source/v1", "Unsupported source index")
        code = text(source.get("document"), "document", 100)
        require(code not in sources, f"Duplicate source: {code}")
        verified_file(source["source"])
        verified_file(source["text"])
        sources[code] = source

    def evidence(refs, owner=None):
        require(isinstance(refs, list) and refs, "Evidence references required")
        for ref in refs:
            require(isinstance(ref, dict), "Evidence must be an object")
            code = ref.get("code")
            require(code in sources and (owner is None or code == owner), "Unknown or wrong proposal reference")
            locator = text(ref.get("locator"), "locator")
            quote = text(ref.get("quote"), "quote")
            blocks = [b for b in sources[code]["blocks"] if b["locator"] == locator]
            require(len(blocks) == 1 and quote in blocks[0]["text"], f"Quote not found at {code} {locator}")

    def claim(value, owner=None):
        require(isinstance(value, dict), "Claim must contain text and evidence")
        text(value.get("text"), "claim", 500)
        evidence(value.get("evidence"), owner)

    def validate_table(table, owner=None):
        require(isinstance(table, dict), "Table must be an object")
        headers, rows = table.get("headers"), table.get("rows")
        require(isinstance(headers, list) and 1 <= len(headers) <= 6, "Use 1–6 table columns")
        require(isinstance(rows, list) and 1 <= len(rows) <= 100, "Use 1–100 table rows")
        for row in [headers, *rows]:
            require(isinstance(row, list) and len(row) == len(headers), "Inconsistent table columns")
            for cell in row:
                text(cell, "table cell", 100)
        evidence(table.get("evidence"), owner)

    proposals = data.get("proposals")
    require(isinstance(proposals, list) and proposals, "At least one analyzed proposal required")
    seen = set()
    for card in proposals:
        require(isinstance(card, dict), "Proposal card must be an object")
        code = card.get("code")
        require(code in sources and code not in seen, f"Unknown or duplicate proposal: {code}")
        seen.add(code)
        for field in ("title", "company", "kind", "spec"):
            text(card.get(field), field, 200)
        require(isinstance(card.get("clauses"), list), "clauses must be a list")
        for clause in card["clauses"]:
            text(clause, "clause", 60)
        claim(card.get("summary"), code)
        points = card.get("points")
        require(isinstance(points, list) and 1 <= len(points) <= 12, "Use 1–12 points per proposal")
        for point in points:
            require(isinstance(point, dict), "Point must be an object")
            text(point.get("title"), "point title", 100)
            require(isinstance(point.get("claims"), list) and 1 <= len(point["claims"]) <= 6, "Use 1–6 claims per point")
            for item in point["claims"]:
                claim(item, code)
            image = point.get("image")
            if image is not None:
                require(isinstance(image, dict), "Image reference must be an object")
                matches = [i for i in sources[code]["images"] if i["id"] == image.get("id")]
                require(len(matches) == 1, "Image must belong to the proposal source index")
                original = matches[0]
                verified_file(original)
                require(image.get("locator") in original["locators"], "Image locator does not match source")
                text(image.get("caption"), "image caption", 160)
                require(image.get("readability") in {"readable", "unreadable"}, "Inspect image and specify readability")
                if image.get("preview"):
                    verified_file(image["preview"])
                if image["readability"] == "readable":
                    path = verified_file(image.get("preview") or original)
                    require(path.suffix.lower() in {".png", ".jpg", ".jpeg"}, "Create and inspect PNG/JPEG preview first; retain original")
                    with Image.open(path) as raster:
                        raster.verify()
            table = point.get("table")
            if table is not None:
                validate_table(table, code)
        require(isinstance(card.get("watch", []), list), "watch must be a list")
        for item in card.get("watch", []):
            claim(item, code)
        require(isinstance(card.get("dataChecks", []), list), "dataChecks must be a list")
        for check in card.get("dataChecks", []):
            require(isinstance(check, dict), "Data check must be an object")
            for field in ("value", "unit", "condition"):
                text(check.get(field), field, 200)
            require(check.get("status") in {"verified", "unverified", "conflict"}, "Invalid data check status")
            evidence(check.get("evidence"), code)
            if check["status"] == "verified":
                require(any(check["value"] in e["quote"] for e in check["evidence"]), "Verified value absent from quoted source")
    require(isinstance(data.get("failures", []), list), "failures must be a list")
    for failure in data.get("failures", []):
        require(isinstance(failure, dict), "Failure must be an object")
        code = text(failure.get("code"), "failed document", 100)
        require(code not in seen, f"Duplicate or both failed and analyzed: {code}")
        text(failure.get("reason"), "failure reason", 500)
        seen.add(code)
    require(set(sources) <= seen, "Source index omitted: supply a card or explicit failure")
    require(isinstance(data.get("synthesis"), list), "synthesis must be a list")
    if len(proposals) > 1:
        require(bool(data["synthesis"]), "Batch deck needs source-linked synthesis")
    for item in data["synthesis"]:
        claim(item)
        require(all(e["code"] in {p["code"] for p in proposals} for e in item["evidence"]), "Synthesis cites an unanalyzed proposal")
        if item.get("table") is not None:
            validate_table(item["table"])
            require(all(e["code"] in {p["code"] for p in proposals} for e in item["table"]["evidence"]), "Comparison table cites an unanalyzed proposal")
    return sources


def audit_boxes(boxes):
    errors = []
    for box in boxes:
        vals = [box[k] for k in ("x", "y", "w", "h")]
        if not all(isinstance(v, (int, float)) and math.isfinite(v) for v in vals):
            errors.append(f"Non-finite geometry: {box['label']}")
            continue
        x, y, w, h = vals
        if min(w, h) <= 0 or x < 0 or y < 0 or x + w > W + .001 or y + h > H + .001:
            errors.append(f"Outside slide: {box['label']}")
        if box.get("zone") == "body" and (y < 1.55 or y + h > 6.4):
            errors.append(f"Outside body area: {box['label']}")
    for i, a in enumerate(boxes):
        for b in boxes[i + 1:]:
            if a["slide"] == b["slide"] and min(a["x"] + a["w"], b["x"] + b["w"]) - max(a["x"], b["x"]) > .02 and min(a["y"] + a["h"], b["y"] + b["h"]) - max(a["y"], b["y"]) > .02:
                errors.append(f"Overlap: {a['label']} / {b['label']}")
    return errors


def wrapped_lines(value, width, size):
    """Conservative estimate only; actual renderer inspection remains required."""
    capacity = max(1, (width * 72 - 8) / size)
    result = 0
    for line in value.split("\n"):
        units = sum(1 if unicodedata.east_asian_width(c) in "WF" else .6 for c in line)
        result += max(1, math.ceil(units / capacity))
    return result


class Deck:
    def __init__(self, font="Noto Sans CJK SC", accent="C00000"):
        require(re.fullmatch(r"[0-9A-Fa-f]{6}", accent) is not None, "accent must be six hex digits")
        self.pres = Presentation()
        self.pres.slide_width, self.pres.slide_height = Inches(W), Inches(H)
        self.font, self.accent, self.boxes = font, accent, []

    def register(self, slide, x, y, w, h, label, zone="body"):
        self.boxes.append(dict(slide=list(self.pres.slides).index(slide), x=x, y=y, w=w, h=h, label=label, zone=zone))

    def style(self, paragraph, value, size, color="27333A", bold=False, highlight=False):
        paragraph.font.name, paragraph.font.size = self.font, Pt(size)
        paragraph.font.color.rgb, paragraph.font.bold = RGBColor.from_string(color), bold
        paragraph.space_after, paragraph.line_spacing = Pt(0), 1.15
        for part in re.split(r"(\d+(?:\.\d+)?%?)", value) if highlight else [value]:
            run = paragraph.add_run()
            run.text = part
            if highlight and re.fullmatch(r"\d+(?:\.\d+)?%?", part):
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(self.accent)

    def textbox(self, slide, value, x, y, w, h, size=16, color="27333A", bold=False, zone="body"):
        require(wrapped_lines(value, w, size) * size * 1.15 + 6 <= h * 72, f"Text needs more room: {value[:60]}")
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.04)
        frame.margin_top = frame.margin_bottom = Inches(.02)
        for n, line in enumerate(value.split("\n")):
            self.style(frame.paragraphs[0] if n == 0 else frame.add_paragraph(), line, size, color, bold)
        self.register(slide, x, y, w, h, value[:40], zone)

    def page(self, title, subtitle="", notes=""):
        slide = self.pres.slides.add_slide(self.pres.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        self.textbox(slide, title, .4, .2, 12.5, .72, 22, self.accent, True, "header")
        if subtitle:
            self.textbox(slide, subtitle, .4, 1.0, 12.5, .42, 12, zone="header")
        slide.notes_slide.notes_text_frame.text = notes
        return slide

    def image(self, slide, path, caption, width=12.4):
        with Image.open(path) as image:
            iw, ih = image.size
        scale = min(width / iw, 4.0 / ih)
        w, h = iw * scale, ih * scale
        x, y = .45 + (width - w) / 2, 1.65 + (4.0 - h) / 2
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        self.register(slide, x, y, w, h, caption)
        self.textbox(slide, caption, .45, 5.85, width, .5, 12)

    def table(self, slide, headers, rows):
        count = len(headers)
        width = 12.4 / count
        heights = [max(wrapped_lines(cell, width, 13) for cell in row) * 13 * 1.15 / 72 + .16 for row in [headers, *rows]]
        height = sum(heights)
        require(height <= 4.7, "Table needs more room; shorten cells or split the table")
        shape = slide.shapes.add_table(len(rows) + 1, count, Inches(.45), Inches(1.65), Inches(12.4), Inches(height))
        for r, values in enumerate([headers, *rows]):
            shape.table.rows[r].height = Inches(heights[r])
            for c, value in enumerate(values):
                cell = shape.table.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(.04)
                cell.margin_top = cell.margin_bottom = Inches(.03)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("2B3A42" if r == 0 else "FFFFFF")
                self.style(cell.text_frame.paragraphs[0], value, 13, "FFFFFF" if r == 0 else "27333A", r == 0, r > 0)
        self.register(slide, .45, 1.65, 12.4, height, "native table")


def references(refs):
    return "\n".join(f"{e['code']} · {e['locator']}\n{e['quote']}" for e in refs)


def build(data, output, font="Noto Sans CJK SC", accent="C00000"):
    sources = validate_cards(data)
    require(not Path(output).exists(), "Use a new output directory; existing results are not overwritten")
    deck = Deck(font, accent)
    warnings = ["文件和原文引用检查不等于结论正确；需检查图表、数字口径和实际排版。"]
    for code, source in sources.items():
        warnings.extend(f"{code}: {warning}" for warning in source.get("warnings", []))
    cover = deck.page(data["title"], data["meeting"])
    codes = [p["code"] for p in data["proposals"]]
    deck.textbox(cover, f"已分析 {len(codes)} 份提案\n公司提案不等于会议已采纳的结论。", .5, 2, 12.2, 1.3, 20)
    if data.get("failures"):
        deck.textbox(cover, f"另有 {len(data['failures'])} 份未完成，见“未完成资料”。", .5, 4, 12.2, .65, 18)
    for card in data["proposals"]:
        code = card["code"]
        subtitle = f"{code} ｜ {card['company']} ｜ {card['kind']} ｜ {card['spec']}"
        notes = references(card["summary"]["evidence"])
        for check in card.get("dataChecks", []):
            notes += f"\n数字使用条件：{check['value']} {check['unit']}；{check['condition']}；{check['status']}\n" + references(check["evidence"])
        slide = deck.page(card["title"], subtitle, notes)
        deck.textbox(slide, card["summary"]["text"], .45, 1.7, 12.4, 2.5, 20)
        deck.textbox(slide, "涉及条款：" + ("、".join(card["clauses"]) or "原文未注明"), .45, 4.6, 12.4, 1.3, 16)
        for point in card["points"]:
            claims = point["claims"]
            consumed = 0
            if point.get("image"):
                image = point["image"]
                original = next(i for i in sources[code]["images"] if i["id"] == image["id"])
                paired = image["readability"] == "readable" and all(wrapped_lines(c["text"], 4.7, 16) * 16 * 1.15 + 6 <= 2.1 * 72 for c in claims[:2]) and wrapped_lines(image["caption"], 7.4, 12) * 12 * 1.15 + 6 <= .5 * 72
                notes = f"{code} · {image['locator']}\n{image['caption']}"
                if paired:
                    notes += "\n" + references([e for c in claims[:2] for e in c["evidence"]])
                slide = deck.page(point["title"], subtitle, notes)
                if image["readability"] == "readable":
                    deck.image(slide, Path((image.get("preview") or original)["path"]), image["caption"], 7.4 if paired else 12.4)
                    if paired:
                        for n, item in enumerate(claims[:2]):
                            deck.textbox(slide, item["text"], 8.15, 1.7 + n * 2.3, 4.7, 2.1, 16)
                            consumed += 1
                else:
                    warning = f"{code}：原图无法辨认，未推测图中内容。"
                    warnings.append(warning)
                    deck.textbox(slide, warning, .45, 2, 12.4, 2, 20)
            # Pair short explanations with their image. Longer text gets separate
            # pages instead of shrinking the font or silently dropping claims.
            for offset in range(consumed, len(claims), 2):
                batch = claims[offset:offset + 2]
                refs = [e for item in batch for e in item["evidence"]]
                slide = deck.page(point["title"], subtitle, references(refs))
                for n, item in enumerate(batch):
                    deck.textbox(slide, item["text"], .45, 1.7 + n * 2.3, 12.4, 2.1, 18)
            if point.get("table"):
                table = point["table"]
                for offset in range(0, len(table["rows"]), 6):
                    slide = deck.page(point["title"], subtitle, references(table["evidence"]))
                    deck.table(slide, table["headers"], table["rows"][offset:offset + 6])
        for item in card.get("watch", []):
            slide = deck.page("需要关注", subtitle, references(item["evidence"]))
            deck.textbox(slide, item["text"], .45, 1.7, 12.4, 4.5, 18)
        for check in card.get("dataChecks", []):
            if check["status"] != "verified":
                warning = f"{code}：{check['value']} {check['unit']}（{check['condition']}）尚未确认或存在口径冲突。"
                warnings.append(warning)
                slide = deck.page("数字使用时请注意", subtitle, references(check["evidence"]))
                deck.textbox(slide, warning, .45, 1.7, 12.4, 4.5, 18)
    for item in data["synthesis"]:
        slide = deck.page("提案对比", data["meeting"], references(item["evidence"]))
        deck.textbox(slide, item["text"], .45, 1.7, 12.4, 4.5, 18)
        if item.get("table"):
            table = item["table"]
            for offset in range(0, len(table["rows"]), 6):
                slide = deck.page("提案对比", data["meeting"], references(table["evidence"]))
                deck.table(slide, table["headers"], table["rows"][offset:offset + 6])
    for failure in data.get("failures", []):
        slide = deck.page("未完成资料", failure["code"])
        deck.textbox(slide, failure["reason"], .45, 1.7, 12.4, 4.5, 18)
    total = len(deck.pres.slides)
    for n, slide in enumerate(deck.pres.slides, 1):
        deck.textbox(slide, f"{n} / {total} · 原文出处见本页备注", .45, 6.85, 12.4, .35, 10, zone="footer")
    errors = audit_boxes(deck.boxes)
    require(not errors, "Layout check failed: " + "; ".join(errors))
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    pptx = output / "proposals.pptx"
    deck.pres.save(pptx)
    require(len(Presentation(pptx).slides) == total, "Saved slide count differs")
    summary = {"schema": "3gpp-proposal-deck/v1", "status": "draft", "deck": {"path": str(pptx.resolve()), "sha256": digest(pptx)},
               "slides": total, "proposals": codes, "failures": data.get("failures", []), "geometry": "passed",
               "visualReview": "not_reviewed", "warnings": warnings}
    (output / "cards.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output / "build-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Source-linked cards JSON")
    parser.add_argument("--output", help="New output directory")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--font", default="Noto Sans CJK SC")
    parser.add_argument("--accent", default="C00000")
    args = parser.parse_args()
    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    if args.validate_only:
        validate_cards(data)
        print("Source files, card structure and quoted locations checked; meaning not verified.")
    else:
        require(args.output, "--output is required")
        print(json.dumps(build(data, args.output, args.font, args.accent), ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except (ValueError, KeyError, TypeError, OSError) as exc:
        raise SystemExit(f"PPT generation failed: {exc}") from exc
